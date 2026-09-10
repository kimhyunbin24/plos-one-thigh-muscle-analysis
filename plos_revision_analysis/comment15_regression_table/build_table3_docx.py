from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ANALYSIS_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ANALYSIS_ROOT / "outputs"
INPUT_JSON = OUTPUT_DIR / "comment15_regression_full.json"
OUTPUT_DOCX = OUTPUT_DIR / "Table3_regression_results.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_properties = cell._tc.get_or_add_tcPr()
    shading = tc_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_properties = cell._tc.get_or_add_tcPr()
    borders = tc_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 70, start: int = 70, bottom: int = 70, end: int = 70) -> None:
    tc = cell._tc
    tc_properties = tc.get_or_add_tcPr()
    margins = tc_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_properties.append(margins)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_properties.append(header)


def set_run_font(run, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")


def p_text(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def build_document() -> None:
    records = json.loads(INPUT_JSON.read_text(encoding="utf-8"))
    records = [
        row
        for row in records
        if row["model"] == "Model 3 adjusted with BMI" and row["predictor"] != "Intercept"
    ]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(16)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(11)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Table 3 Candidate Multivariable Regression Results")
    set_run_font(title_run, 16, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    intro_run = intro.add_run(
        "Unstandardized coefficients from the BMI-included adjusted model reconstructed from "
        "the available revision code. Author review is required before manuscript insertion."
    )
    set_run_font(intro_run, 9)

    outcomes = []
    for row in records:
        if row["outcome"] not in outcomes:
            outcomes.append(row["outcome"])

    headers = ["Predictor", "Reference", "Beta", "SE", "95% CI", "p", "n"]
    widths = [2.35, 1.65, 0.85, 0.85, 1.75, 0.75, 0.7]

    for outcome_index, outcome in enumerate(outcomes):
        if outcome_index in (2, 4):
            doc.add_page_break()
        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(7 if outcome_index else 2)
        heading.paragraph_format.space_after = Pt(3)
        sample = next(row for row in records if row["outcome"] == outcome)
        heading_run = heading.add_run(f"{outcome}  {sample['outcome_unit']}")
        set_run_font(heading_run, 11, bold=True)

        outcome_rows = [row for row in records if row["outcome"] == outcome]
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        set_repeat_table_header(table.rows[0])

        for column_index, header_text in enumerate(headers):
            cell = table.rows[0].cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "1F4E78")
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(header_text)
            set_run_font(run, 8, bold=True, color="FFFFFF")

        for row_index, result in enumerate(outcome_rows):
            row = table.add_row()
            values = [
                result["predictor"],
                result["reference_category"],
                f"{float(result['beta']):.4f}",
                f"{float(result['standard_error']):.4f}",
                f"{float(result['ci_lower']):.4f} to {float(result['ci_upper']):.4f}",
                p_text(float(result["p_value"])),
                str(result["sample_size"]),
            ]
            for column_index, value in enumerate(values):
                cell = row.cells[column_index]
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_borders(cell)
                set_cell_margins(cell)
                if row_index % 2 == 1:
                    set_cell_shading(cell, "EDF3F8")
                paragraph = cell.paragraphs[0]
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column_index < 2 else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(str(value))
                set_run_font(run, 7.7)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(7)
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run(
        "Model covariates: age, sex, fracture type, Koval grade, FIM grade, CCI, anesthesia type, "
        "surgery type, and BMI. Female sex, femoral neck fracture, general anesthesia, and "
        "arthroplasty are reference categories. Koval and FIM were entered as reversed ordinal "
        "grades. Fracture side was not included in the available code. Fat outcomes use fat/pure "
        "multiplied by 10, matching the current script but not a conventional percentage scale. "
        "The Total thigh outcome uses the source variable femoral_total_volume; confirm its "
        "anatomical meaning before publication."
    )
    set_run_font(note_run, 8)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()

