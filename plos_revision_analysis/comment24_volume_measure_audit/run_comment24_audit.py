from __future__ import annotations

import math
import statistics
import sys
from pathlib import Path

import numpy as np
from docx import Document
from openpyxl import load_workbook


ANALYSIS_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = ANALYSIS_ROOT.parent
if str(ANALYSIS_ROOT / "shared") not in sys.path:
    sys.path.insert(0, str(ANALYSIS_ROOT / "shared"))

from analysis_core import (  # noqa: E402
    OUTPUT_DIR,
    excel_path,
    load_analysis_rows,
    manuscript_path,
    values,
    write_csv,
)


def source_line(path: Path, needle: str) -> str:
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if needle in line:
            return f"{path.name}:{number}"
    return f"{path.name}: pattern not found"


def audit_rows() -> list[dict[str, object]]:
    table2 = PROJECT_ROOT / "table2_print.py"
    revision = PROJECT_ROOT / "revision_adjusted_analysis.py"
    manuscript = Document(manuscript_path())
    paragraphs = {index: paragraph.text for index, paragraph in enumerate(manuscript.paragraphs)}

    return [
        {
            "analysis": "Table 2 descriptive volume",
            "actual_variable_used": "Excel total, pure, and fat volume columns divided by 1000",
            "unit": "cm3",
            "calculation": "source mm3 / 1000",
            "evidence": source_line(table2, "mean_sd_by_header(rows, headers, f\"{prefix}_total_volume_mm3\", 1000)"),
            "status": "Verified from code and reproduced output",
        },
        {
            "analysis": "Table 2 fat measure",
            "actual_variable_used": "Excel fat_percentage column divided by 10",
            "unit": "Labeled percent in manuscript",
            "calculation": "Excel fat/pure*100 result / 10",
            "evidence": source_line(table2, "mean_sd_by_header(rows, headers, f\"{prefix}_fat_percentage\", 10)"),
            "status": "INCONSISTENCY FOUND: division by 10 makes the reported value one tenth of the source percentage",
        },
        {
            "analysis": "Table 2 total thigh label",
            "actual_variable_used": "femoral_total_volume_mm3",
            "unit": "cm3",
            "calculation": "femoral_total_volume_mm3 / 1000",
            "evidence": source_line(table2, 'muscle_row("Total thigh muscle volume", rows, headers, "femoral")'),
            "status": "INCONSISTENCY FOUND: variable identity does not demonstrate that this is the sum of thigh muscle groups",
        },
        {
            "analysis": "Anatomical-group comparison",
            "actual_variable_used": "Derived group fat-to-pure ratio",
            "unit": "Manuscript scale",
            "calculation": "normalized fat / normalized pure * 10",
            "evidence": source_line(revision, "return fat / pure * 10.0"),
            "status": "Verified in current revision code; mathematically not a percent scale",
        },
        {
            "analysis": "Age-group muscle analysis",
            "actual_variable_used": "Not verifiable",
            "unit": "Not verifiable",
            "calculation": "Manuscript states ANOVA across age categories, but no executable age-group analysis is present",
            "evidence": "manuscript.docx paragraph 46; no matching project script",
            "status": "INCONSISTENCY FOUND: reported method cannot be reproduced from supplied code",
        },
        {
            "analysis": "Sex-related muscle analysis",
            "actual_variable_used": "Not verifiable for muscle outcomes",
            "unit": "Not verifiable",
            "calculation": "No muscle-volume sex comparison script was found",
            "evidence": "manuscript.docx Results; no matching project script",
            "status": "Not reproducible from supplied code",
        },
        {
            "analysis": "Multivariable regression volume outcomes",
            "actual_variable_used": "femoral_total_volume, gluteal_total, and anterior_Total from SAV multiplied by 10",
            "unit": "cm3/m2",
            "calculation": "upstream mm3/cm2 normalized value * 10",
            "evidence": source_line(revision, '("Total thigh normalized volume, cm3/m2", values(rows, "femoral_total_volume") * 10.0)'),
            "status": "Verified in current revision code; only three volume outcomes are implemented",
        },
        {
            "analysis": "Multivariable regression fat outcomes",
            "actual_variable_used": "gluteal and anterior fat-to-pure ratios",
            "unit": "Manuscript scale",
            "calculation": "normalized fat / normalized pure * 10",
            "evidence": source_line(revision, "group_fat_percent(rows, \"gluteal_pure\", \"gluteal_fat\")"),
            "status": "INCONSISTENCY FOUND: scale differs from source percentage by factor 10",
        },
        {
            "analysis": "Height normalization described in manuscript",
            "actual_variable_used": "Volume divided by height squared upstream before SAV export",
            "unit": "cm3/m2 after unit conversion",
            "calculation": "Excel mm3 / height_cm2, then multiply by 10",
            "evidence": f"manuscript.docx paragraph 44; code conversion at {source_line(revision, 'total_v = values(rows, total) * 10.0')}",
            "status": "Formula verified numerically; upstream normalization code is not present",
        },
        {
            "analysis": "Manuscript statistical model specification",
            "actual_variable_used": "Current code uses continuous age, omits fracture side, and provides BMI included and excluded models",
            "unit": "Not applicable",
            "calculation": "Current code differs from manuscript paragraph 47",
            "evidence": "manuscript.docx paragraph 47 and revision_adjusted_analysis.py design_matrix",
            "status": "INCONSISTENCY FOUND: model covariates and primary-model designation differ",
        },
    ]


def scale_check() -> tuple[list[dict[str, object]], dict[str, float]]:
    workbook = load_workbook(excel_path(), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    data_rows = [
        row
        for row in sheet.iter_rows(min_row=3, values_only=True)
        if len(row) > 142
        and str(row[1]).strip() == "수술전"
        and isinstance(row[26], (int, float))
        and not isinstance(row[26], bool)
    ]

    group_columns = [
        ("Anterior", 27, 28, 29, 30),
        ("Medial", 51, 52, 53, 54),
        ("Gluteal", 87, 88, 89, 90),
        ("Posterior", 103, 104, 105, 106),
        ("Others", 127, 128, 129, 130),
    ]
    checks = []
    for group, total_col, fat_col, pure_col, percent_col in group_columns:
        total_i, fat_i, pure_i, percent_i = [column - 1 for column in (total_col, fat_col, pure_col, percent_col)]
        raw_percent = [float(row[percent_i]) for row in data_rows]
        fat_to_pure = [float(row[fat_i]) / float(row[pure_i]) * 100.0 for row in data_rows]
        fat_to_total = [float(row[fat_i]) / float(row[total_i]) * 100.0 for row in data_rows]
        checks.append(
            {
                "group": group,
                "n": len(raw_percent),
                "excel_source_percentage_mean": statistics.mean(raw_percent),
                "recomputed_fat_divided_by_pure_times_100_mean": statistics.mean(fat_to_pure),
                "recomputed_fat_divided_by_total_times_100_mean": statistics.mean(fat_to_total),
                "manuscript_reported_scale_mean_excel_divided_by_10": statistics.mean(raw_percent) / 10.0,
                "maximum_absolute_source_formula_difference": max(
                    abs(a - b) for a, b in zip(raw_percent, fat_to_pure)
                ),
                "scale_factor_source_to_manuscript": 10.0,
            }
        )

    sav_rows = load_analysis_rows()
    sav_normalized = values(sav_rows, "anterior_Total")
    excel_normalized = np.asarray(
        [float(row[26]) / float(row[142]) ** 2 for row in data_rows], dtype=float
    )
    normalization = {
        "n": float(len(sav_rows)),
        "mae_mm3_per_cm2": float(np.mean(np.abs(sav_normalized - excel_normalized))),
        "max_abs_difference_mm3_per_cm2": float(np.max(np.abs(sav_normalized - excel_normalized))),
        "correlation": float(np.corrcoef(sav_normalized, excel_normalized)[0, 1]),
        "femoral_absolute_mean_cm3": statistics.mean(float(row[130]) for row in data_rows) / 1000.0,
        "four_primary_groups_sum_mean_cm3": statistics.mean(
            float(row[26]) + float(row[50]) + float(row[86]) + float(row[102])
            for row in data_rows
        ) / 1000.0,
    }
    workbook.close()
    return checks, normalization


def write_report(audit: list[dict[str, object]], checks: list[dict[str, object]], normalization: dict[str, float]) -> None:
    lines = [
        "# Comment 24 volume measure audit",
        "",
        "| Analysis | Actual variable used | Unit | Evidence | Status |",
        "|---|---|---|---|---|",
    ]
    for row in audit:
        lines.append(
            f"| {row['analysis']} | {row['actual_variable_used']} | {row['unit']} | "
            f"{row['evidence']} | {row['status']} |"
        )

    lines.extend(
        [
            "",
            "## Height-normalization check",
            "",
            "The SAV anterior-group normalized volume agrees with Excel anterior total volume "
            "divided by height in centimeters squared. The mean absolute difference is "
            f"{normalization['mae_mm3_per_cm2']:.6f} mm3/cm2, the maximum absolute difference "
            f"is {normalization['max_abs_difference_mm3_per_cm2']:.6f}, and the correlation is "
            f"{normalization['correlation']:.9f}. The small discrepancy reflects the two-decimal "
            "rounding in the SAV file. Multiplication by 10 converts mm3/cm2 to cm3/m2.",
            "",
            "The Table 2 row labeled Total thigh muscle volume uses femoral_total_volume_mm3. "
            f"Its mean is {normalization['femoral_absolute_mean_cm3']:.2f} cm3, whereas the sum "
            "of the anterior, medial, gluteal, and posterior group volumes has a mean of "
            f"{normalization['four_primary_groups_sum_mean_cm3']:.2f} cm3. The femoral variable "
            "therefore is not the sum of those four groups; its anatomical definition must be confirmed.",
            "",
            "## Fat percentage scale check",
            "",
            "The Excel percentage columns equal fat volume divided by pure volume multiplied by "
            "100. The Table 2 code divides those values by 10 before printing them. The reported "
            "means are therefore one tenth of the source percentage values.",
            "",
            "| Group | Source fat to pure percent | Manuscript scale | Fat to total percent |",
            "|---|---:|---:|---:|",
        ]
    )
    for row in checks:
        lines.append(
            f"| {row['group']} | {float(row['excel_source_percentage_mean']):.3f}% | "
            f"{float(row['manuscript_reported_scale_mean_excel_divided_by_10']):.3f}% | "
            f"{float(row['recomputed_fat_divided_by_total_times_100_mean']):.3f}% |"
        )

    lines.extend(
        [
            "",
            "## Potential inconsistencies requiring author review",
            "",
            "1. Table 2 uses absolute volume in cm3 even though the Methods state that volumes "
            "were normalized to height squared.",
            "2. The printed fat percentage is one tenth of the Excel fat-to-pure percentage. "
            "The denominator and scale must be defined and corrected consistently.",
            "3. No executable age-group ANOVA or muscle-volume sex comparison was supplied.",
            "4. The reconstructable regression code uses continuous age, omits fracture side, "
            "and uses selected outcomes, whereas the manuscript describes age-group and fracture-side adjustment.",
            "5. The SAV variables named fat_percentage are all zero and were not used by the "
            "revision script. Fat ratios were recalculated from normalized fat and pure volumes.",
            "6. The label Total thigh muscle volume is assigned to femoral_total_volume_mm3, "
            "which is not the sum of the four main anatomical groups.",
        ]
    )
    (OUTPUT_DIR / "comment24_volume_measure_audit.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )


def main() -> None:
    audit = audit_rows()
    checks, normalization = scale_check()
    write_csv(
        OUTPUT_DIR / "comment24_volume_measure_audit.csv",
        audit,
        ["analysis", "actual_variable_used", "unit", "calculation", "evidence", "status"],
    )
    write_csv(
        OUTPUT_DIR / "comment24_fat_percentage_scale_check.csv",
        checks,
        [
            "group",
            "n",
            "excel_source_percentage_mean",
            "recomputed_fat_divided_by_pure_times_100_mean",
            "recomputed_fat_divided_by_total_times_100_mean",
            "manuscript_reported_scale_mean_excel_divided_by_10",
            "maximum_absolute_source_formula_difference",
            "scale_factor_source_to_manuscript",
        ],
    )
    write_report(audit, checks, normalization)
    print(f"Audit rows written: {len(audit)}")
    print(f"Fat scale check groups: {len(checks)}")
    print(OUTPUT_DIR / "comment24_volume_measure_audit.md")


if __name__ == "__main__":
    main()

