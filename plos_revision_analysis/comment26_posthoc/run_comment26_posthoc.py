from __future__ import annotations

import csv
import math
from itertools import combinations
from pathlib import Path

import numpy as np
from docx import Document

from plos_revision_analysis.shared.analysis_core import (
    OUTPUT_DIR,
    complete_mask,
    load_analysis_rows,
    regularized_beta,
    student_t_critical_975,
    student_t_two_sided_p,
    values,
)


COMPARTMENTS = {
    "Anterior": "anterior_Total",
    "Medial": "medial_total",
    "Posterior": "posterior_total",
    "Gluteal": "gluteal_total",
}

MUSCLES = {
    "Sartorius": "sartorius_total_volume",
    "Rectus femoris": "rectus_femoris_total_volume",
    "Vastus lateralis": "vastus_lateralis_total_volume",
    "Vastus intermedius": "vastus_intermedius_total_volume",
    "Vastus medialis": "vastus_medialis_total_volume",
    "Adductor longus": "adductor_longus_total_volume",
    "Adductor brevis": "adductor_brevis_total_volume",
    "Adductor magnus": "adductor_magnus_total_volume",
    "Gracilis": "gracilis_total_volume",
    "Pectineus": "pectineus_total_volume",
    "Semitendinosus": "semitendinosus_total_volume",
    "Semimembranosus": "semimembranosus_total_volume",
    "Biceps femoris": "biceps_femoris_total_volume",
    "Gluteus maximus": "gluteus_maximus_total_volume",
    "Gluteus medius": "gluteus_medius_total_volume",
    "Gluteus minimus": "gluteus_minimus_total_volume",
    "Tensor fasciae latae": "tensor_fascia_latae_total_volume",
    "Piriformis": "piriformis_total_volume",
    "Obturator internus": "obturator_internus_total_volume",
    "Obturator externus": "obturator_externus_total_volume",
    "Quadratus femoris": "quadratus_femoris_total_volume",
}


def matrix(rows: list[dict[str, object]], specs: dict[str, str]) -> tuple[list[str], np.ndarray]:
    names = list(specs)
    columns = [values(rows, specs[name]) * 10.0 for name in names]
    data = np.column_stack(columns)
    return names, data[np.all(np.isfinite(data), axis=1)]


def repeated_measures_anova(data: np.ndarray) -> dict[str, float | int]:
    n, k = data.shape
    grand = float(data.mean())
    condition_means = data.mean(axis=0)
    subject_means = data.mean(axis=1)
    ss_total = float(((data - grand) ** 2).sum())
    ss_condition = float(n * ((condition_means - grand) ** 2).sum())
    ss_subject = float(k * ((subject_means - grand) ** 2).sum())
    ss_error = ss_total - ss_condition - ss_subject
    df1 = k - 1
    df2 = (n - 1) * (k - 1)
    f_value = (ss_condition / df1) / (ss_error / df2)
    x = (df1 * f_value) / (df1 * f_value + df2)
    p_value = 1.0 - regularized_beta(df1 / 2.0, df2 / 2.0, x)
    covariance = np.cov(data, rowvar=False, ddof=1)
    centering = np.eye(k) - np.ones((k, k)) / k
    centered_covariance = centering @ covariance @ centering
    epsilon = float(
        np.trace(centered_covariance) ** 2
        / ((k - 1) * np.trace(centered_covariance @ centered_covariance))
    )
    gg_df1 = epsilon * df1
    gg_df2 = epsilon * df2
    gg_x = (gg_df1 * f_value) / (gg_df1 * f_value + gg_df2)
    gg_p = 1.0 - regularized_beta(gg_df1 / 2.0, gg_df2 / 2.0, gg_x)
    return {
        "n": n, "conditions": k, "df1": df1, "df2": df2, "f": f_value, "p": p_value,
        "gg_epsilon": epsilon, "gg_df1": gg_df1, "gg_df2": gg_df2, "gg_p": gg_p,
    }


def holm_adjust(p_values: list[float]) -> list[float]:
    order = sorted(range(len(p_values)), key=p_values.__getitem__)
    adjusted = [math.nan] * len(p_values)
    running = 0.0
    m = len(p_values)
    for rank, index in enumerate(order):
        running = max(running, min(1.0, (m - rank) * p_values[index]))
        adjusted[index] = running
    return adjusted


def paired_posthoc(names: list[str], data: np.ndarray, family: str) -> list[dict[str, object]]:
    rows = []
    for left, right in combinations(range(len(names)), 2):
        difference = data[:, left] - data[:, right]
        n = len(difference)
        mean_difference = float(difference.mean())
        standard_error = float(difference.std(ddof=1) / math.sqrt(n))
        t_value = mean_difference / standard_error
        p_value = student_t_two_sided_p(t_value, n - 1)
        critical = student_t_critical_975(n - 1)
        rows.append({
            "family": family,
            "comparison": f"{names[left]} vs {names[right]}",
            "n": n,
            "mean_difference_cm3_m2": mean_difference,
            "ci_lower": mean_difference - critical * standard_error,
            "ci_upper": mean_difference + critical * standard_error,
            "t_value": t_value,
            "df": n - 1,
            "p_raw": p_value,
        })
    adjusted = holm_adjust([float(row["p_raw"]) for row in rows])
    for row, p_adjusted in zip(rows, adjusted):
        row["p_holm"] = p_adjusted
        row["significant_holm_0_05"] = p_adjusted < 0.05
    return rows


def p_text(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_word(path: Path, omnibus: dict[str, float | int], compartment_rows: list[dict[str, object]], muscle_rows: list[dict[str, object]]) -> None:
    document = Document()
    document.add_heading("Repeated-measures muscle-volume comparisons", level=1)
    document.add_paragraph(
        "Outcome: height-squared normalized total muscle volume (cm3/m2). "
        "The four-compartment analysis is primary. Individual-muscle comparisons form a separate exploratory multiplicity family."
    )
    document.add_paragraph(
        f"Repeated-measures ANOVA: F({omnibus['df1']}, {omnibus['df2']}) = {omnibus['f']:.3f}, "
        f"p {p_text(float(omnibus['p']))}, N = {omnibus['n']}. "
        f"Greenhouse-Geisser epsilon = {omnibus['gg_epsilon']:.3f}; corrected "
        f"F({omnibus['gg_df1']:.2f}, {omnibus['gg_df2']:.2f}) = {omnibus['f']:.3f}, "
        f"p {p_text(float(omnibus['gg_p']))}."
    )
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Comparison", "N", "Mean difference", "95% CI", "Raw p", "Holm p"]):
        cell.text = label
    for result in compartment_rows:
        cells = table.add_row().cells
        cells[0].text = str(result["comparison"])
        cells[1].text = str(result["n"])
        cells[2].text = f"{result['mean_difference_cm3_m2']:.2f}"
        cells[3].text = f"{result['ci_lower']:.2f} to {result['ci_upper']:.2f}"
        cells[4].text = p_text(float(result["p_raw"]))
        cells[5].text = p_text(float(result["p_holm"]))
    document.add_paragraph(
        f"Individual-muscle exploratory family: {len(muscle_rows)} paired comparisons; "
        f"{sum(bool(row['significant_holm_0_05']) for row in muscle_rows)} remained significant after Holm correction. "
        "Complete results are provided in the companion CSV file."
    )
    document.save(path)


def main() -> None:
    rows = load_analysis_rows()
    compartment_names, compartment_data = matrix(rows, COMPARTMENTS)
    muscle_names, muscle_data = matrix(rows, MUSCLES)
    omnibus = repeated_measures_anova(compartment_data)
    compartment_results = paired_posthoc(compartment_names, compartment_data, "Primary compartments")
    muscle_results = paired_posthoc(muscle_names, muscle_data, "Exploratory individual muscles")

    output_dir = OUTPUT_DIR / "comment26_posthoc"
    output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(output_dir / "comment26_compartment_posthoc.csv", compartment_results)
    write_csv(output_dir / "comment26_individual_muscle_posthoc.csv", muscle_results)
    write_word(output_dir / "Comment26_posthoc_results.docx", omnibus, compartment_results, muscle_results)

    response = (
        "We revised the analysis to account for the within-participant structure of the anatomical measurements. "
        f"A one-way repeated-measures ANOVA showed an overall compartment difference "
        f"(Greenhouse-Geisser-corrected F({omnibus['gg_df1']:.2f}, {omnibus['gg_df2']:.2f}) = "
        f"{omnibus['f']:.3f}, p {p_text(float(omnibus['gg_p']))}; epsilon = {omnibus['gg_epsilon']:.3f}). "
        "Post hoc comparisons were performed using paired t-tests with Holm adjustment across the six primary "
        "compartment comparisons. Individual-muscle comparisons were treated as a separate exploratory family "
        "and were Holm-adjusted within that family. Exact p-values and 95% confidence intervals have been added "
        "to the revised table. IMAT was not included in this analysis because its final operational definition "
        "and scaling remain under audit."
    )
    (output_dir / "comment26_response_draft.txt").write_text(response, encoding="utf-8")

    print(f"N = {omnibus['n']}")
    print(f"Repeated-measures ANOVA: F({omnibus['df1']}, {omnibus['df2']}) = {omnibus['f']:.3f}, p {p_text(float(omnibus['p']))}")
    print(f"Greenhouse-Geisser: epsilon = {omnibus['gg_epsilon']:.3f}, "
          f"F({omnibus['gg_df1']:.2f}, {omnibus['gg_df2']:.2f}) = {omnibus['f']:.3f}, "
          f"p {p_text(float(omnibus['gg_p']))}")
    for result in compartment_results:
        print(f"{result['comparison']}: mean difference {result['mean_difference_cm3_m2']:.2f}, "
              f"95% CI {result['ci_lower']:.2f} to {result['ci_upper']:.2f}, "
              f"raw p {p_text(float(result['p_raw']))}, Holm p {p_text(float(result['p_holm']))}")
    print(output_dir / "Comment26_posthoc_results.docx")


if __name__ == "__main__":
    main()

