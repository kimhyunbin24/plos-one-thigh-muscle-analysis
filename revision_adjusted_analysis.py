from __future__ import annotations

import math
import os
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from table1_print import read_spss_sav


DATA_DIR = Path(os.environ.get("CT_VOLUME_DATA_DIR", Path(__file__).with_name("data")))
SAV_PATH = Path(
    os.environ.get("CT_VOLUME_SAV_PATH", DATA_DIR / "analysis_data.sav")
)


GROUPS = [
    ("Anterior", "anterior_Total", "anterior_pure", "anterior_fat"),
    ("Medial", "medial_total", "medial_pure", "medial_fat"),
    ("Posterior", "posterior_total", "posterior_pure", "posterior_fat"),
    ("Gluteal", "gluteal_total", "gluteal_pure", "gluteal_fat"),
    ("Others", "other_total", "other_pure", "other_fat"),
    ("Total thigh", "femoral_total_volume", "femoral_pure_volume", "femoral_fat_volume"),
]

OUTPUT_NORMALIZED = Path(__file__).with_name("revision_normalized_volume_table.docx")
OUTPUT_PAIRWISE = Path(__file__).with_name("revision_pairwise_comparison_table.docx")
OUTPUT_REGRESSION = Path(__file__).with_name("revision_adjusted_regression_table.docx")
OUTPUT_VIF = Path(__file__).with_name("revision_vif_table.docx")


def normal_cdf(x: float) -> float:
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def normal_p_from_z(z: float) -> float:
    return 2.0 * (1.0 - normal_cdf(abs(z)))


def as_float(value: object) -> float:
    if isinstance(value, (int, float)) and not math.isnan(float(value)):
        return float(value)
    return math.nan


def load_rows() -> list[dict[str, float | str]]:
    return read_spss_sav(SAV_PATH)


def values(rows: list[dict[str, float | str]], name: str) -> np.ndarray:
    return np.array([as_float(row.get(name)) for row in rows], dtype=float)


def complete_mask(*arrays: np.ndarray) -> np.ndarray:
    mask = np.ones(len(arrays[0]), dtype=bool)
    for arr in arrays:
        mask &= np.isfinite(arr)
    return mask


def mean_sd(arr: np.ndarray) -> str:
    arr = arr[np.isfinite(arr)]
    return f"{arr.mean():.2f}+/-{arr.std(ddof=1):.2f}"


def fmt_p(p_value: float) -> str:
    if not math.isfinite(p_value):
        return ""
    if p_value < 0.001:
        return "<0.001"
    return f"{p_value:.3f}"


def bh_fdr(p_values: list[float]) -> list[float]:
    n = len(p_values)
    order = sorted(range(n), key=lambda i: p_values[i])
    adjusted = [1.0] * n
    running = 1.0
    for rank_from_end, i in enumerate(reversed(order), start=1):
        rank = n - rank_from_end + 1
        running = min(running, p_values[i] * n / rank)
        adjusted[i] = min(running, 1.0)
    return adjusted


def design_matrix(rows: list[dict[str, float | str]], include_bmi: bool = False) -> tuple[np.ndarray, list[str]]:
    age = values(rows, "age")
    sex_male = (values(rows, "sex") == 1).astype(float)
    fx_intertrochanter = (values(rows, "fx_type") == 1).astype(float)
    koval_grade = 8.0 - values(rows, "KOVAL")
    fim_grade = 8.0 - values(rows, "FIM")
    cci = values(rows, "CCI")
    spinal = (values(rows, "anesthesia") == 2).astype(float)
    epidural = (values(rows, "anesthesia") == 3).astype(float)
    internal_fixation = (values(rows, "surgery_type") == 2).astype(float)

    columns = [
        np.ones(len(rows)),
        age,
        sex_male,
        fx_intertrochanter,
        koval_grade,
        fim_grade,
        cci,
        spinal,
        epidural,
        internal_fixation,
    ]
    names = [
        "Intercept",
        "Age",
        "Male sex",
        "Intertrochanteric fracture",
        "Koval grade",
        "FIM grade",
        "CCI",
        "Spinal anesthesia",
        "Epidural anesthesia",
        "Internal fixation",
    ]

    if include_bmi:
        columns.append(values(rows, "BMI"))
        names.append("BMI")

    x = np.column_stack(columns)
    return x, names


def ols(y: np.ndarray, x: np.ndarray, names: list[str]) -> list[dict[str, float | str]]:
    mask = complete_mask(y, *[x[:, i] for i in range(x.shape[1])])
    y = y[mask]
    x = x[mask]

    beta = np.linalg.lstsq(x, y, rcond=None)[0]
    resid = y - x @ beta
    df = x.shape[0] - x.shape[1]
    sigma2 = float(resid.T @ resid / df)
    cov = sigma2 * np.linalg.inv(x.T @ x)
    se = np.sqrt(np.diag(cov))
    z = beta / se

    out = []
    for name, b, s, zi in zip(names, beta, se, z):
        out.append(
            {
                "term": name,
                "beta": float(b),
                "se": float(s),
                "ci_low": float(b - 1.96 * s),
                "ci_high": float(b + 1.96 * s),
                "p": float(normal_p_from_z(float(zi))),
                "n": int(x.shape[0]),
            }
        )
    return out


def vif_table(x: np.ndarray, names: list[str]) -> list[tuple[str, float]]:
    out = []
    for i in range(1, x.shape[1]):
        y = x[:, i]
        others = np.delete(x, i, axis=1)
        mask = complete_mask(y, *[others[:, j] for j in range(others.shape[1])])
        y2 = y[mask]
        o2 = others[mask]
        beta = np.linalg.lstsq(o2, y2, rcond=None)[0]
        resid = y2 - o2 @ beta
        sse = float(resid.T @ resid)
        sst = float(((y2 - y2.mean()) ** 2).sum())
        r2 = 1.0 - sse / sst if sst else 0.0
        vif = 1.0 / (1.0 - r2) if r2 < 1 else math.inf
        out.append((names[i], vif))
    return out


def print_height2_descriptives(rows: list[dict[str, float | str]]) -> None:
    print("\n=== Height-squared normalized muscle volume descriptive table ===")
    print("Unit: cm3/m2. SAV stores mm3/cm2, so values are multiplied by 10.")
    print("Group\tTotal volume\tPure volume\tFat volume")
    for row in normalized_volume_rows(rows):
        print("\t".join(row))


def normalized_volume_rows(rows: list[dict[str, float | str]]) -> list[tuple[str, str, str, str]]:
    out = []
    for label, total, pure, fat in GROUPS:
        total_v = values(rows, total) * 10.0
        pure_v = values(rows, pure) * 10.0
        fat_v = values(rows, fat) * 10.0
        out.append((label, mean_sd(total_v), mean_sd(pure_v), mean_sd(fat_v)))
    return out


def group_fat_percent(rows: list[dict[str, float | str]], pure_name: str, fat_name: str) -> np.ndarray:
    pure = values(rows, pure_name)
    fat = values(rows, fat_name)
    # Same scale as manuscript Table 2 fat percentage: Excel percentage / 10.
    return fat / pure * 10.0


def print_pairwise_group_comparisons(rows: list[dict[str, float | str]]) -> None:
    print("\n=== Pairwise anatomical-group fat percentage comparisons ===")
    print("Outcome: fat volume / pure volume * 10, matching manuscript Table 2 scale.")

    print("Comparison\tn\tMean diff\t95% CI\tp\tp_Bonferroni\tp_FDR")
    for row in pairwise_group_rows(rows):
        print("\t".join(row))


def pairwise_group_rows(rows: list[dict[str, float | str]]) -> list[tuple[str, str, str, str, str, str, str]]:
    group_values = {
        label: group_fat_percent(rows, pure, fat)
        for label, _total, pure, fat in GROUPS
        if label != "Total thigh"
    }

    raw_results = []
    labels = list(group_values)
    for i in range(len(labels)):
        for j in range(i + 1, len(labels)):
            a = group_values[labels[i]]
            b = group_values[labels[j]]
            mask = complete_mask(a, b)
            diff = a[mask] - b[mask]
            n = len(diff)
            mean_diff = float(diff.mean())
            se = float(diff.std(ddof=1) / math.sqrt(n))
            z = mean_diff / se
            p = normal_p_from_z(z)
            raw_results.append(
                {
                    "comparison": f"{labels[i]} - {labels[j]}",
                    "n": n,
                    "mean_diff": mean_diff,
                    "ci_low": mean_diff - 1.96 * se,
                    "ci_high": mean_diff + 1.96 * se,
                    "p": p,
                }
            )

    fdr = bh_fdr([r["p"] for r in raw_results])
    bonf = [min(r["p"] * len(raw_results), 1.0) for r in raw_results]

    out = []
    for result, p_bonf, p_fdr in zip(raw_results, bonf, fdr):
        out.append(
            (
                str(result["comparison"]),
                str(result["n"]),
                f"{result['mean_diff']:.3f}",
                f"{result['ci_low']:.3f} to {result['ci_high']:.3f}",
                fmt_p(float(result["p"])),
                fmt_p(float(p_bonf)),
                fmt_p(float(p_fdr)),
            )
        )
    return out


def print_regression_block(
    rows: list[dict[str, float | str]],
    outcome_label: str,
    y: np.ndarray,
) -> None:
    print(f"\n=== Regression outcome: {outcome_label} ===")

    x_basic = np.column_stack([np.ones(len(rows)), values(rows, "age"), (values(rows, "sex") == 1).astype(float)])
    basic_names = ["Intercept", "Age", "Male sex"]
    basic = ols(y, x_basic, basic_names)
    print("\nModel 1: unadjusted/basic predictors")
    print("Term\tBeta\t95% CI\tp\tn")
    for r in basic[1:]:
        print(f"{r['term']}\t{r['beta']:.4f}\t{r['ci_low']:.4f} to {r['ci_high']:.4f}\t{fmt_p(float(r['p']))}\t{r['n']}")

    x_adj, adj_names = design_matrix(rows, include_bmi=False)
    adjusted = ols(y, x_adj, adj_names)
    print("\nModel 2: adjusted model without BMI")
    print("Term\tBeta\t95% CI\tp\tn")
    for r in adjusted[1:]:
        print(f"{r['term']}\t{r['beta']:.4f}\t{r['ci_low']:.4f} to {r['ci_high']:.4f}\t{fmt_p(float(r['p']))}\t{r['n']}")

    x_bmi, bmi_names = design_matrix(rows, include_bmi=True)
    adjusted_bmi = ols(y, x_bmi, bmi_names)
    print("\nModel 3: sensitivity model with BMI")
    print("Term\tBeta\t95% CI\tp\tn")
    for r in adjusted_bmi[1:]:
        print(f"{r['term']}\t{r['beta']:.4f}\t{r['ci_low']:.4f} to {r['ci_high']:.4f}\t{fmt_p(float(r['p']))}\t{r['n']}")


def regression_rows(rows: list[dict[str, float | str]]) -> list[tuple[str, str, str, str, str, str, str]]:
    out = []
    outcomes = [
        ("Total thigh normalized volume, cm3/m2", values(rows, "femoral_total_volume") * 10.0),
        ("Gluteal normalized volume, cm3/m2", values(rows, "gluteal_total") * 10.0),
        ("Anterior normalized volume, cm3/m2", values(rows, "anterior_Total") * 10.0),
        ("Gluteal fat percentage", group_fat_percent(rows, "gluteal_pure", "gluteal_fat")),
        ("Anterior fat percentage", group_fat_percent(rows, "anterior_pure", "anterior_fat")),
    ]
    for outcome_label, y in outcomes:
        models = [
            ("Model 1 basic", *ols(y, np.column_stack([np.ones(len(rows)), values(rows, "age"), (values(rows, "sex") == 1).astype(float)]), ["Intercept", "Age", "Male sex"])[1:]),
        ]
        for r in models[0][1:]:
            out.append(
                (
                    outcome_label,
                    models[0][0],
                    str(r["term"]),
                    f"{float(r['beta']):.4f}",
                    f"{float(r['ci_low']):.4f} to {float(r['ci_high']):.4f}",
                    fmt_p(float(r["p"])),
                    str(r["n"]),
                )
            )

        x_adj, adj_names = design_matrix(rows, include_bmi=False)
        for r in ols(y, x_adj, adj_names)[1:]:
            out.append(
                (
                    outcome_label,
                    "Model 2 adjusted without BMI",
                    str(r["term"]),
                    f"{float(r['beta']):.4f}",
                    f"{float(r['ci_low']):.4f} to {float(r['ci_high']):.4f}",
                    fmt_p(float(r["p"])),
                    str(r["n"]),
                )
            )

        x_bmi, bmi_names = design_matrix(rows, include_bmi=True)
        for r in ols(y, x_bmi, bmi_names)[1:]:
            out.append(
                (
                    outcome_label,
                    "Model 3 sensitivity with BMI",
                    str(r["term"]),
                    f"{float(r['beta']):.4f}",
                    f"{float(r['ci_low']):.4f} to {float(r['ci_high']):.4f}",
                    fmt_p(float(r["p"])),
                    str(r["n"]),
                )
            )
    return out


def print_vif(rows: list[dict[str, float | str]]) -> None:
    print("\n=== Multicollinearity diagnostics: VIF ===")
    for include_bmi in [False, True]:
        x, names = design_matrix(rows, include_bmi=include_bmi)
        print("\nAdjusted model " + ("with BMI" if include_bmi else "without BMI"))
        print("Predictor\tVIF")
        for name, vif in vif_table(x, names):
            print(f"{name}\t{vif:.2f}")


def vif_rows_for_docx(rows: list[dict[str, float | str]]) -> list[tuple[str, str, str]]:
    out = []
    for include_bmi in [False, True]:
        x, names = design_matrix(rows, include_bmi=include_bmi)
        model = "Adjusted with BMI" if include_bmi else "Adjusted without BMI"
        for name, vif in vif_table(x, names):
            out.append((model, name, f"{vif:.2f}"))
    return out


def make_docx_table(
    title: str,
    headers: tuple[str, ...],
    rows: list[tuple[str, ...]],
    output_path: Path,
    note: str = "",
    landscape: bool = False,
) -> None:
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
    else:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(12)

    if note:
        note_p = doc.add_paragraph()
        note_p.paragraph_format.space_after = Pt(6)
        note_run = note_p.add_run(note)
        note_run.font.name = "Calibri"
        note_run.font.size = Pt(8)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for col_i, header in enumerate(headers):
        cell = table.rows[0].cells[col_i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(8)

    for row_i, row_values in enumerate(rows, start=1):
        for col_i, value in enumerate(row_values):
            cell = table.rows[row_i].cells[col_i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.name = "Calibri"
            run.font.size = Pt(7.5 if landscape else 8)

    doc.save(output_path)


def save_result_docx(rows: list[dict[str, float | str]]) -> None:
    make_docx_table(
        "Revision Table. Height-squared normalized muscle volume",
        ("Group", "Total volume, cm3/m2", "Pure volume, cm3/m2", "Fat volume, cm3/m2"),
        normalized_volume_rows(rows),
        OUTPUT_NORMALIZED,
        "Values are mean +/- SD. SAV values were converted to cm3/m2 by multiplying by 10.",
    )
    make_docx_table(
        "Revision Table. Pairwise anatomical-group fat percentage comparisons",
        ("Comparison", "n", "Mean diff", "95% CI", "p", "p Bonferroni", "p FDR"),
        pairwise_group_rows(rows),
        OUTPUT_PAIRWISE,
        "Outcome is fat volume / pure volume * 10. p-values use normal approximation; Bonferroni and FDR corrections are shown.",
        landscape=True,
    )
    make_docx_table(
        "Revision Table. Regression analyses for normalized volume and fat percentage",
        ("Outcome", "Model", "Term", "Beta", "95% CI", "p", "n"),
        regression_rows(rows),
        OUTPUT_REGRESSION,
        "Model 2 excludes BMI as the primary adjusted model. Model 3 includes BMI as sensitivity analysis.",
        landscape=True,
    )
    make_docx_table(
        "Revision Table. Variance inflation factor diagnostics",
        ("Model", "Predictor", "VIF"),
        vif_rows_for_docx(rows),
        OUTPUT_VIF,
        "VIF was calculated for adjusted models with and without BMI.",
    )


def main() -> None:
    rows = load_rows()
    print(f"Analysis dataset n = {len(rows)}")

    print_height2_descriptives(rows)
    print_pairwise_group_comparisons(rows)

    outcomes = [
        ("Total thigh normalized volume, cm3/m2", values(rows, "femoral_total_volume") * 10.0),
        ("Gluteal normalized volume, cm3/m2", values(rows, "gluteal_total") * 10.0),
        ("Anterior normalized volume, cm3/m2", values(rows, "anterior_Total") * 10.0),
        ("Gluteal fat percentage", group_fat_percent(rows, "gluteal_pure", "gluteal_fat")),
        ("Anterior fat percentage", group_fat_percent(rows, "anterior_pure", "anterior_fat")),
    ]
    for label, y in outcomes:
        print_regression_block(rows, label, y)

    print_vif(rows)
    save_result_docx(rows)
    print("\nSaved Word tables:")
    print(OUTPUT_NORMALIZED)
    print(OUTPUT_PAIRWISE)
    print(OUTPUT_REGRESSION)
    print(OUTPUT_VIF)


if __name__ == "__main__":
    main()

